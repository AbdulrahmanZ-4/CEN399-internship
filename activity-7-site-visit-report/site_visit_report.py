import os
import shutil
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


# ==============================
# Folder setup
# ==============================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
REPORTS_DIR = os.path.join(BASE_DIR, "reports")
PHOTOS_DIR = os.path.join(BASE_DIR, "site_photos")

os.makedirs(REPORTS_DIR, exist_ok=True)
os.makedirs(PHOTOS_DIR, exist_ok=True)


# ==============================
# Helper functions
# ==============================

def sanitize_filename(name):
    """
    Removes characters that are not allowed in Windows file names.
    """
    invalid_chars = '<>:"/\\|?*'
    for char in invalid_chars:
        name = name.replace(char, "_")
    return name.strip()


def get_text(widget):
    """
    Gets text from Entry or Text widgets.
    """
    if isinstance(widget, tk.Text):
        return widget.get("1.0", tk.END).strip()
    return widget.get().strip()


def add_bullet_lines(document, text):
    """
    Adds multiline text as bullet points in the Word report.
    """
    lines = [line.strip() for line in text.split("\n") if line.strip()]

    if not lines:
        document.add_paragraph("N/A")
        return

    for line in lines:
        if line.startswith("-"):
            line = line[1:].strip()
        document.add_paragraph(line, style="List Bullet")


def set_document_style(document):
    """
    Sets the default font style for the report.
    """
    styles = document.styles
    normal_style = styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)


def add_info_table(document, data):
    """
    Adds the basic site visit information in a table.
    """
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    info_rows = [
        ("Company", data["company"]),
        ("Project Name", data["project_name"]),
        ("Site Location", data["site_location"]),
        ("Visit Date", data["visit_date"]),
        ("Inspector / Engineer", data["inspector"]),
        ("Contractor", data["contractor"]),
        ("Consultant", data["consultant"]),
        ("Weather", data["weather"]),
    ]

    for label, value in info_rows:
        row_cells = table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = value if value else "N/A"

    document.add_paragraph()


def create_report(data, photos):
    """
    Creates the Word site visit report.
    """
    document = Document()
    set_document_style(document)

    # Title
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("SITE VISIT REPORT")
    title_run.bold = True
    title_run.font.size = Pt(18)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(data["company"])
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(12)

    document.add_paragraph()

    # Section 1
    document.add_heading("1. General Information", level=1)
    add_info_table(document, data)

    # Section 2
    document.add_heading("2. Work Progress Observed", level=1)
    add_bullet_lines(document, data["work_progress"])

    # Section 3
    document.add_heading("3. Safety Observations", level=1)
    add_bullet_lines(document, data["safety_notes"])

    # Section 4
    document.add_heading("4. Quality Observations", level=1)
    add_bullet_lines(document, data["quality_notes"])

    # Section 5
    document.add_heading("5. Issues / Concerns", level=1)
    add_bullet_lines(document, data["issues"])

    # Section 6
    document.add_heading("6. Recommended Next Actions", level=1)
    add_bullet_lines(document, data["next_actions"])

    # Section 7
    document.add_heading("7. Site Photos", level=1)

    if not photos:
        document.add_paragraph("No site photos were attached.")
    else:
        for index, photo in enumerate(photos, start=1):
            photo_path = photo["path"]
            caption = photo["caption"]

            document.add_paragraph(f"Photo {index}: {caption if caption else 'Site photo'}")

            try:
                document.add_picture(photo_path, width=Inches(5.5))
            except Exception:
                document.add_paragraph("[Image could not be inserted]")

            document.add_paragraph()

    # Section 8
    document.add_heading("8. Prepared By", level=1)
    document.add_paragraph(f"Prepared by: {data['inspector'] if data['inspector'] else 'N/A'}")
    document.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    # Save file
    project_name = sanitize_filename(data["project_name"]) or "Site_Visit"
    visit_date = sanitize_filename(data["visit_date"]) or datetime.now().strftime("%Y-%m-%d")

    file_name = f"Site_Visit_Report_{project_name}_{visit_date}.docx"
    output_path = os.path.join(REPORTS_DIR, file_name)

    document.save(output_path)

    return output_path


# ==============================
# GUI Application
# ==============================

class SiteVisitReportApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Site Visit Report Generator")
        self.root.geometry("900x750")

        self.fields = {}
        self.photos = []

        self.create_gui()

    def create_gui(self):
        # Main title
        title_label = tk.Label(
            self.root,
            text="Site Visit Report Generator",
            font=("Arial", 18, "bold")
        )
        title_label.pack(pady=10)

        # Scrollable frame
        container = tk.Frame(self.root)
        container.pack(fill="both", expand=True)

        canvas = tk.Canvas(container)
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        self.scroll_frame = tk.Frame(canvas)

        self.scroll_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=self.scroll_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        row = 0

        # Basic fields
        self.add_entry("Company", "company", row, default="Scope Consulting Engineers - سكوب للاستشارات الهندسية")
        row += 1

        self.add_entry("Project Name", "project_name", row)
        row += 1

        self.add_entry("Site Location", "site_location", row)
        row += 1

        self.add_entry("Visit Date", "visit_date", row, default=datetime.now().strftime("%Y-%m-%d"))
        row += 1

        self.add_entry("Inspector / Engineer", "inspector", row)
        row += 1

        self.add_entry("Contractor", "contractor", row)
        row += 1

        self.add_entry("Consultant", "consultant", row, default="Scope Consulting Engineers")
        row += 1

        self.add_entry("Weather", "weather", row, default="Sunny")
        row += 1

        # Long text fields
        self.add_textbox("Work Progress Observed", "work_progress", row)
        row += 1

        self.add_textbox("Safety Observations", "safety_notes", row)
        row += 1

        self.add_textbox("Quality Observations", "quality_notes", row)
        row += 1

        self.add_textbox("Issues / Concerns", "issues", row)
        row += 1

        self.add_textbox("Recommended Next Actions", "next_actions", row)
        row += 1

        # Photos section
        photos_label = tk.Label(
            self.scroll_frame,
            text="Site Photos",
            font=("Arial", 12, "bold")
        )
        photos_label.grid(row=row, column=0, sticky="w", padx=15, pady=10)

        photo_buttons_frame = tk.Frame(self.scroll_frame)
        photo_buttons_frame.grid(row=row, column=1, sticky="w", padx=15, pady=10)

        add_photo_btn = tk.Button(
            photo_buttons_frame,
            text="Add Photos",
            command=self.add_photos,
            width=15
        )
        add_photo_btn.pack(side="left", padx=5)

        remove_photo_btn = tk.Button(
            photo_buttons_frame,
            text="Remove Selected",
            command=self.remove_selected_photo,
            width=15
        )
        remove_photo_btn.pack(side="left", padx=5)

        row += 1

        self.photos_listbox = tk.Listbox(self.scroll_frame, width=85, height=6)
        self.photos_listbox.grid(row=row, column=0, columnspan=2, padx=15, pady=5)

        row += 1

        # Generate button
        generate_btn = tk.Button(
            self.scroll_frame,
            text="Generate Word Report",
            command=self.generate_report,
            bg="#2E7D32",
            fg="white",
            font=("Arial", 13, "bold"),
            width=25
        )
        generate_btn.grid(row=row, column=0, columnspan=2, pady=25)

    def add_entry(self, label_text, field_key, row, default=""):
        label = tk.Label(
            self.scroll_frame,
            text=label_text,
            font=("Arial", 10, "bold")
        )
        label.grid(row=row, column=0, sticky="w", padx=15, pady=7)

        entry = tk.Entry(self.scroll_frame, width=80)
        entry.grid(row=row, column=1, sticky="w", padx=15, pady=7)

        if default:
            entry.insert(0, default)

        self.fields[field_key] = entry

    def add_textbox(self, label_text, field_key, row):
        label = tk.Label(
            self.scroll_frame,
            text=label_text,
            font=("Arial", 10, "bold")
        )
        label.grid(row=row, column=0, sticky="nw", padx=15, pady=7)

        text_box = tk.Text(self.scroll_frame, width=60, height=5)
        text_box.grid(row=row, column=1, sticky="w", padx=15, pady=7)

        sample_text = "- "
        text_box.insert("1.0", sample_text)

        self.fields[field_key] = text_box

    def add_photos(self):
        file_paths = filedialog.askopenfilenames(
            title="Select Site Photos",
            filetypes=[
                ("Image Files", "*.jpg *.jpeg *.png"),
                ("All Files", "*.*")
            ]
        )

        if not file_paths:
            return

        for file_path in file_paths:
            try:
                file_name = os.path.basename(file_path)
                copied_path = os.path.join(PHOTOS_DIR, file_name)

                # Avoid overwriting files with same name
                if os.path.exists(copied_path):
                    name, ext = os.path.splitext(file_name)
                    new_name = f"{name}_{datetime.now().strftime('%H%M%S')}{ext}"
                    copied_path = os.path.join(PHOTOS_DIR, new_name)

                shutil.copy(file_path, copied_path)

                caption = simpledialog.askstring(
                    "Photo Caption",
                    f"Enter caption for:\n{file_name}"
                )

                self.photos.append({
                    "path": copied_path,
                    "caption": caption if caption else ""
                })

            except Exception as e:
                messagebox.showerror("Photo Error", f"Could not add photo:\n{file_path}\n\n{e}")

        self.update_photos_listbox()

    def remove_selected_photo(self):
        selected = self.photos_listbox.curselection()

        if not selected:
            messagebox.showwarning("No Selection", "Please select a photo to remove.")
            return

        index = selected[0]
        self.photos.pop(index)
        self.update_photos_listbox()

    def update_photos_listbox(self):
        self.photos_listbox.delete(0, tk.END)

        for index, photo in enumerate(self.photos, start=1):
            file_name = os.path.basename(photo["path"])
            caption = photo["caption"] if photo["caption"] else "No caption"
            self.photos_listbox.insert(tk.END, f"{index}. {file_name} - {caption}")

    def collect_data(self):
        data = {}

        for key, widget in self.fields.items():
            data[key] = get_text(widget)

        return data

    def generate_report(self):
        data = self.collect_data()

        if not data["project_name"]:
            messagebox.showwarning("Missing Information", "Please enter the project name.")
            return

        if not data["inspector"]:
            messagebox.showwarning("Missing Information", "Please enter the inspector/engineer name.")
            return

        try:
            output_path = create_report(data, self.photos)

            messagebox.showinfo(
                "Report Generated",
                f"Site visit report created successfully:\n\n{output_path}"
            )

        except Exception as e:
            messagebox.showerror("Error", f"Could not generate report:\n\n{e}")


# ==============================
# Run app
# ==============================

if __name__ == "__main__":
    root = tk.Tk()
    app = SiteVisitReportApp(root)
    root.mainloop()